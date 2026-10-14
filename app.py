import os
import re
import io
import shutil
import datetime
import pyautogui
from PIL import Image
import pandas as pd
from pypdf import PdfReader
from docx import Document
import openpyxl
import requests
from bs4 import BeautifulSoup
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pydub import AudioSegment
from faster_whisper import WhisperModel

app = FastAPI(title="Offline Personal PQM AI Assistant Engine")

# --- CONFIGURATION ---
STORAGE_DIR = "./pqm_vault"  
OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "llama3.2:1b"  
REQUEST_TIMEOUT = 300  # 5 minutes to accommodate CPU inference without timing out

os.makedirs(STORAGE_DIR, exist_ok=True)
os.makedirs(os.path.join(STORAGE_DIR, "audits"), exist_ok=True)

print("Loading offline Whisper STT engine...")
whisper_engine = WhisperModel("base", device="cpu", compute_type="float32")
print("Whisper STT engine loaded successfully.")

class QueryRequest(BaseModel):
    query: str

# --- CORE FILE EXTRACTORS ---

def extract_txt(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()[:8000]
    except Exception as e:
        return f"[Error reading TXT {os.path.basename(file_path)}: {str(e)}]"

def extract_pdf(file_path):
    text = ""
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        text = f"[Error reading PDF {os.path.basename(file_path)}: {str(e)}]"
    return text[:10000]

def extract_docx(file_path):
    try:
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])[:8000]
    except Exception as e:
        return f"[Error reading Word file {os.path.basename(file_path)}: {str(e)}]"

def extract_xlsx(file_path):
    text = f"--- Spreadsheet: {os.path.basename(file_path)} ---\n"
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet in wb.sheetnames:
            text += f"Sheet: {sheet}\n"
            ws = wb[sheet]
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if any(row):  
                    text += " | ".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
                    row_count += 1
                if row_count >= 100:
                    text += "... [Rows truncated for brevity] ...\n"
                    break
    except Exception as e:
        text += f"[Error reading Excel: {str(e)}]\n"
    return text[:8000]

def scrape_url(url):
    try:
        headers = {'User-Agent': 'Mozilla/5.0'}
        r = requests.get(url, headers=headers, timeout=10)
        if r.status_code == 200:
            soup = BeautifulSoup(r.text, 'html.parser')
            for s in soup(["script", "style", "nav", "footer", "header"]):
                s.extract()
            lines = (line.strip() for line in soup.get_text().splitlines())
            return "\n".join(chunk for chunk in lines if chunk)[:8000] 
    except Exception as e:
        return f"[Failed to scrape live URL: {str(e)}]"
    return "[Unable to fetch URL]"

# --- OS AUTOMATION & DATA COMPILATION ---

def take_desktop_screenshot():
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    screenshot_path = os.path.join(STORAGE_DIR, f"screenshot_{timestamp}.png")
    screenshot = pyautogui.screenshot()
    screenshot.save(screenshot_path)
    return screenshot_path

def compile_all_local_context(query_text=""):
    """
    Intelligent context compiler:
    If the user mentions specific file names in their prompt, only those files
    are loaded into the LLM context. Otherwise, reads the entire vault.
    """
    master_context = ""
    query_lower = query_text.lower()
    
    all_vault_files = []
    for root, _, files in os.walk(STORAGE_DIR):
        for file in files:
            all_vault_files.append((root, file))

    matched_files = [
        (r, f) for r, f in all_vault_files 
        if f.lower() in query_lower or os.path.splitext(f.lower())[0] in query_lower
    ]

    target_list = matched_files if matched_files else all_vault_files

    for root, file in target_list:
        path = os.path.join(root, file)
        ext = file.split('.')[-1].lower()
        rel_dir = os.path.relpath(path, STORAGE_DIR)
        
        master_context += f"\n\n=== START OF FILE: {rel_dir} ===\n"
        if ext in ['txt', 'log']:
            master_context += extract_txt(path)
        elif ext == 'pdf':
            master_context += extract_pdf(path)
        elif ext == 'docx':
            master_context += extract_docx(path)
        elif ext in ['xlsx', 'xls']:
            master_context += extract_xlsx(path)
        master_context += f"\n=== END OF FILE: {rel_dir} ===\n"
        
    return master_context

# --- STATIC UI ROUTE ---

@app.get("/")
async def serve_frontend_dashboard():
    if os.path.exists("app.html"):
        return FileResponse("app.html")
    raise HTTPException(status_code=404, detail="app.html file missing from the current working directory.")

# --- VAULT REPOSITORY MANAGEMENT ---

@app.post("/upload-vault")
async def upload_document_to_vault(
    file: UploadFile = File(...),
    subfolder: str = Form("")
):
    clean_subfolder = os.path.normpath(subfolder.strip()).lstrip("/\\")
    target_dir = os.path.join(STORAGE_DIR, clean_subfolder)
    os.makedirs(target_dir, exist_ok=True)
    
    file_path = os.path.join(target_dir, file.filename)
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        
        display_path = os.path.relpath(file_path, STORAGE_DIR)
        return {
            "status": "success",
            "filename": file.filename,
            "saved_rel_path": display_path
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to save document to vault: {str(e)}")

# --- AI WORKSPACE ROUTING ---

@app.post("/chat")
async def process_assistant_query(payload: QueryRequest):
    user_query = payload.query.lower()
    screenshot_triggered = False
    screenshot_file = None
    scraped_context = ""

    url_match = re.search(r'(https?://[^\s]+)', payload.query)
    if url_match:
        target_url = url_match.group(1)
        scraped_context = f"\n=== LIVE SCRAPED CONTEXT FROM {target_url} ===\n" + scrape_url(target_url)

    if any(k in user_query for k in ["screenshot", "show screen", "capture screen"]):
        try:
            screenshot_file = take_desktop_screenshot()
            screenshot_triggered = True
        except Exception as e:
            print(f"Automation execution failed: {e}")

    vault_context = compile_all_local_context(payload.query)
    
    system_prompt = (
        "You are an offline personal Project Quality Management (PQM) assistant operating with zero cloud connectivity.\n"
        "Your task is to review compliance checklists, RFI tracking logs, assessments, timelines, and chats to resolve user queries.\n"
        "Be punchy, clear, and group your findings under explicit logical headings.\n"
        f"Current Time context: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
        f"--- LOCAL DATA CONTEXT STORAGE ---\n{vault_context}\n{scraped_context}\n"
    )
    
    if screenshot_triggered:
        system_prompt += f"\n[SYSTEM ALERT: A desktop screenshot was saved locally at '{screenshot_file}'.]\n"

    ollama_payload = {
        "model": MODEL_NAME,
        "prompt": f"{system_prompt}\nUser Query: {payload.query}\nAssistant Answer:",
        "stream": False,
        "options": {
            "num_ctx": 4096,
            "num_predict": 1024
        }
    }

    try:
        response = requests.post(OLLAMA_URL, json=ollama_payload, timeout=REQUEST_TIMEOUT)
        if response.status_code == 200:
            ai_response = response.json().get("response", "No readable output generated by local AI engine.")
            if screenshot_triggered:
                ai_response = f"📸 **[System Action: Screen Captured]** Saved locally as `{os.path.basename(screenshot_file)}`.\n\n" + ai_response
            return {"response": ai_response}
        else:
            raise HTTPException(status_code=500, detail=f"Ollama returned internal error: {response.text}")
    except requests.exceptions.Timeout:
        raise HTTPException(status_code=504, detail="Inference timed out after 5 minutes. Consider narrowing the query.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to communicate with local model engine: {str(e)}")


@app.post("/chat-audio")
async def process_audio_assistant_query(
    audio_file: UploadFile = File(...)
):
    if not audio_file.filename.lower().endswith('.opus'):
        raise HTTPException(status_code=400, detail="Only WhatsApp audio files (.opus) are allowed.")

    try:
        opus_content = await audio_file.read()
        opus_io = io.BytesIO(opus_content)

        audio_segment = AudioSegment.from_file(opus_io, codec="opus")

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        base_name = os.path.splitext(audio_file.filename)[0]
        mp3_filename = f"converted_{base_name}_{timestamp}.mp3"
        saved_mp3_path = os.path.join(STORAGE_DIR, mp3_filename)

        audio_segment.export(saved_mp3_path, format="mp3")

        segments, info = whisper_engine.transcribe(saved_mp3_path, beam_size=5)
        transcribed_text = " ".join([segment.text for segment in segments]).strip()

        if not transcribed_text:
            raise HTTPException(status_code=400, detail="Could not capture or transcribe speech from recording.")

        vault_context = compile_all_local_context(transcribed_text)
        
        system_prompt = (
            "You are an offline personal Project Quality Management (PQM) assistant operating with zero cloud connectivity.\n"
            "Your task is to review compliance checklists, RFI tracking logs, assessments, timelines, and chats to resolve user queries.\n"
            "Be punchy, clear, and group your findings under explicit logical headings.\n"
            f"Current Time context: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
            f"--- LOCAL DATA CONTEXT STORAGE ---\n{vault_context}\n"
        )

        ollama_payload = {
            "model": MODEL_NAME,
            "prompt": f"{system_prompt}\nUser Audio Query (Transcribed): {transcribed_text}\nAssistant Answer:",
            "stream": False,
            "options": {
                "num_ctx": 4096,
                "num_predict": 1024
            }
        }

        response = requests.post(OLLAMA_URL, json=ollama_payload, timeout=REQUEST_TIMEOUT)
        
        if response.status_code == 200:
            ai_response = response.json().get("response", "No readable output generated by local AI engine.")
            return {
                "response": ai_response,
                "speech_transcript": transcribed_text,
                "audio_saved_as": mp3_filename,
                "detected_language": info.language
            }
        else:
            raise HTTPException(status_code=500, detail=f"Ollama execution failed: {response.text}")

    except requests.exceptions.Timeout:
        raise HTTPException(status_code=504, detail="Audio evaluation timed out. Please try again.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Internal pipeline failure during audio review: {str(e)}")


@app.post("/review-procedure")
async def review_procedure(
    procedure_pdf: UploadFile = File(...),
    project_requirements: str = Form("Check if this procedure complies with quality management standards.")
):
    try:
        pdf_content = await procedure_pdf.read()

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        archived_name = f"{timestamp}_{procedure_pdf.filename}"
        archive_path = os.path.join(STORAGE_DIR, "audits", archived_name)
        with open(archive_path, "wb") as f:
            f.write(pdf_content)

        pdf_file = io.BytesIO(pdf_content)
        reader = PdfReader(pdf_file)

        extracted_text = ""
        for page in reader.pages:
            text = page.extract_text()
            if text:
                extracted_text += text + "\n"

        if not extracted_text.strip():
            raise HTTPException(status_code=400, detail="The uploaded PDF contains no extractable text layer.")

        system_instruction = """
        You are an expert Project Quality Management (PQM) compliance auditor operating offline.
        Your task is to stringently cross-examine the provided operational procedure text against the user's explicit project requirements.
        Format your answer cleanly in Markdown using distinct sections: Summary, Compliant Items, Gaps/Non-Compliance, and Action Items.
        """
        full_prompt = (
            f"{system_instruction}\n\n"
            f"[PROJECT REQUIREMENTS TO CHECK AGAINST]:\n{project_requirements}\n\n"
            f"[EXTRACTED PROCEDURE TEXT]:\n{extracted_text[:8000]}\n\n"
            f"Assistant Audit Report:"
        )

        ollama_payload = {
            "model": MODEL_NAME,
            "prompt": full_prompt,
            "stream": False,
            "options": {
                "num_ctx": 4096,
                "num_predict": 1024
            }
        }

        response = requests.post(OLLAMA_URL, json=ollama_payload, timeout=REQUEST_TIMEOUT)

        if response.status_code == 200:
            ai_response = response.json().get("response", "No evaluation report generated by local engine.")
            return {
                "compliance_report": ai_response,
                "saved_in_vault": f"audits/{archived_name}"
            }
        else:
            raise HTTPException(status_code=500, detail=f"Ollama evaluation failed: {response.text}")

    except requests.exceptions.Timeout:
        raise HTTPException(
            status_code=504, 
            detail="The local AI engine took longer than 5 minutes to audit this document. Consider auditing a shorter section."
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Internal pipeline failure during PDF review: {str(e)}")

# --- DIRECT EXECUTION ENTRYPOINT ---

if __name__ == "__main__":
    import uvicorn
    print("Starting PQM Engine on http://localhost:8000 ...")
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)